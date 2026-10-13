import re
from decimal import Decimal
from html import unescape
from io import BytesIO
from pathlib import Path
from uuid import UUID

from docx import Document
from jinja2 import BaseLoader, Environment
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.documents.context_builder import build_document_context
from app.modules.documents.models import DocumentTemplate, DocumentType, GeneratedDocument
from app.modules.documents.numbering import allocate_document_number

jinja_env = Environment(loader=BaseLoader(), autoescape=True)


def render_template_html(body_html: str, context: dict[str, object]) -> str:
    template = jinja_env.from_string(body_html)
    return template.render(**context)


def html_to_docx_bytes(html: str) -> bytes:
    document = Document()
    text = unescape(re.sub(r"<[^>]+>", "\n", html))
    for line in [chunk.strip() for chunk in text.splitlines() if chunk.strip()]:
        document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def generate_document(
    session: AsyncSession,
    *,
    doc_type: DocumentType,
    title: str,
    template: DocumentTemplate,
    lead_id: UUID | None,
    project_id: UUID | None,
    calculation_id: UUID | None,
    amount: Decimal | None,
    vat_rate: Decimal | None,
    vat_included: bool,
) -> GeneratedDocument:
    number, year = await allocate_document_number(session, doc_type)
    resolved_amount = amount
    context = await build_document_context(
        session,
        doc_type=doc_type,
        lead_id=lead_id,
        project_id=project_id,
        calculation_id=calculation_id,
        amount=resolved_amount,
        vat_rate=vat_rate,
        vat_included=vat_included,
        number=number,
    )
    rendered = render_template_html(template.body_html, context)
    if resolved_amount is None and context.get("amount"):
        resolved_amount = Decimal(str(context["amount"]))
    document = GeneratedDocument(
        doc_type=doc_type,
        title=title,
        number=number,
        number_year=year,
        lead_id=lead_id,
        project_id=project_id,
        calculation_id=calculation_id,
        template_id=template.id,
        amount=resolved_amount,
        vat_rate=vat_rate,
        vat_included=vat_included,
        context_data=context,
        rendered_html=rendered,
    )
    session.add(document)
    await session.flush()
    document.root_id = document.id
    return document


def store_scan_bytes(storage_root: Path, document_id: UUID, filename: str, payload: bytes) -> str:
    folder = storage_root / "scans"
    folder.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^A-Za-z0-9_.-]", "_", filename)
    path = folder / f"{document_id}_{safe_name}"
    path.write_bytes(payload)
    return str(path.relative_to(storage_root))
